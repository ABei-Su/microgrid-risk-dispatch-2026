"""Validate official-format result workbooks against frozen seed-1 intervals."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CASES = {
    "result2.xlsx": ("q2_interval_results.csv", False),
    "result3.xlsx": ("q3_interval_results.csv", True),
    "result4-2.xlsx": ("q4_2_interval_results.csv", False),
    "result4-3.xlsx": ("q4_3_interval_results.csv", True),
}


def values(ws, row1: int, row2: int, col1: int, col2: int) -> np.ndarray:
    return np.array([[ws.cell(row, col).value for col in range(col1, col2 + 1)] for row in range(row1, row2 + 1)], dtype=float)


def check_result1(path: Path) -> list[dict]:
    source = pd.read_csv(ROOT / "result" / "Q1" / "Q1_final_trajectories.csv")
    source = source.loc[source["seed"].eq(1)].sort_values("slot")
    workbook = load_workbook(path, data_only=True)
    actual = values(workbook["计划购电量"], 2, 145, 2, 2).ravel()
    block_charge = np.array([workbook["充放电量"].cell(row, 2).value for row in range(2, 8)], dtype=float)
    block_discharge = np.array([workbook["充放电量"].cell(row, 3).value for row in range(2, 8)], dtype=float)
    expected_charge = np.array([source.iloc[block * 24 : (block + 1) * 24]["charge_kwh"].sum() for block in range(6)])
    expected_discharge = np.array([source.iloc[block * 24 : (block + 1) * 24]["discharge_kwh"].sum() for block in range(6)])
    return [
        {"file": path.name, "check": "official_sheet_names", "pass": workbook.sheetnames == ["计划购电量", "充放电量"]},
        {"file": path.name, "check": "all_144_plan_values", "pass": bool(np.allclose(actual, source["grid_kwh"], atol=5.1e-7, rtol=0))},
        {"file": path.name, "check": "plan_total", "pass": bool(abs(actual.sum() - 59482.69899598835) < 1e-4)},
        {"file": path.name, "check": "six_charge_blocks", "pass": bool(np.allclose(block_charge, expected_charge, atol=5.1e-7, rtol=0))},
        {"file": path.name, "check": "six_discharge_blocks", "pass": bool(np.allclose(block_discharge, expected_discharge, atol=5.1e-7, rtol=0))},
        {"file": path.name, "check": "initial_and_final_soc", "pass": abs(float(workbook["充放电量"]["E2"].value) - float(source.iloc[0]["soc_start_kwh"])) < 1e-6 and abs(float(workbook["充放电量"]["E3"].value) - float(source.iloc[-1]["soc_end_kwh"])) < 1e-6},
    ]


def check_annual(path: Path, source_name: str, adjusted: bool) -> list[dict]:
    source = pd.read_csv(ROOT / "figure" / "data" / source_name)
    days = [day.sort_values("interval_index") for _, day in source.groupby("date", sort=True)]
    workbook = load_workbook(path, data_only=True)
    expected_sheets = ["计划购电量"] + (["调整购电量"] if adjusted else []) + ["充放电量", "紧急购电量"]
    checks = [{"file": path.name, "check": "official_sheet_names", "pass": workbook.sheetnames == expected_sheets}]
    for sheet_name, column in [("计划购电量", "base_grid_kwh")] + ([("调整购电量", "effective_grid_kwh")] if adjusted else []):
        actual = values(workbook[sheet_name], 2, 335, 2, 145)
        expected = np.array([day[column].to_numpy(float) for day in days])
        totals = values(workbook[sheet_name], 2, 335, 146, 146).ravel()
        fees = values(workbook[sheet_name], 2, 335, 147, 147).ravel()
        dates = [pd.Timestamp(workbook[sheet_name].cell(row, 1).value).strftime("%Y-%m-%d") for row in range(2, 336)]
        if sheet_name == "调整购电量":
            expected_fees = []
            for day in days:
                price = day["price_yuan_per_kwh"].to_numpy(float)
                base = day["base_grid_kwh"].to_numpy(float)
                effective = day["effective_grid_kwh"].to_numpy(float)
                expected_fees.append(price @ base + 1.5 * price @ np.maximum(effective - base, 0) - 0.5 * price @ np.maximum(base - effective, 0))
        else:
            expected_fees = [day["price_yuan_per_kwh"].to_numpy(float) @ day[column].to_numpy(float) for day in days]
        checks.extend([
            {"file": path.name, "check": f"{sheet_name}_all_48096_values", "pass": actual.shape == (334, 144) and bool(np.allclose(actual, expected, atol=5.1e-7, rtol=0))},
            {"file": path.name, "check": f"{sheet_name}_daily_totals", "pass": bool(np.allclose(totals, expected.sum(axis=1), atol=1e-5, rtol=0))},
            {"file": path.name, "check": f"{sheet_name}_daily_fees", "pass": bool(np.allclose(fees, expected_fees, atol=5.1e-5, rtol=0))},
            {"file": path.name, "check": f"{sheet_name}_334_dates", "pass": dates == sorted(source["date"].unique().tolist())},
        ])
    charge = workbook["充放电量"]
    checks.append({"file": path.name, "check": "charge_discharge_334x6_rows", "pass": charge.max_row == 2005})
    charge_ok = True
    for index, day in enumerate(days):
        start = 2 + index * 6
        actual_charge = np.array([charge.cell(start + block, 3).value for block in range(6)], dtype=float)
        actual_discharge = np.array([charge.cell(start + block, 4).value for block in range(6)], dtype=float)
        expected_charge = np.array([day.iloc[block * 24 : (block + 1) * 24]["charge_kwh"].sum() for block in range(6)])
        expected_discharge = np.array([day.iloc[block * 24 : (block + 1) * 24]["discharge_kwh"].sum() for block in range(6)])
        charge_ok &= bool(np.allclose(actual_charge, expected_charge, atol=5.1e-7, rtol=0))
        charge_ok &= bool(np.allclose(actual_discharge, expected_discharge, atol=5.1e-7, rtol=0))
        charge_ok &= abs(float(charge.cell(start, 6).value) - float(day.iloc[0]["soc_start_kwh"])) < 1e-6
        charge_ok &= abs(float(charge.cell(start + 1, 6).value) - float(day.iloc[-1]["soc_end_kwh"])) < 1e-6
    checks.append({"file": path.name, "check": "all_charge_discharge_and_soc_values", "pass": charge_ok})
    emergency_total = sum(float(row[2] or 0) for row in workbook["紧急购电量"].iter_rows(min_row=2, values_only=True))
    checks.append({"file": path.name, "check": "emergency_total", "pass": abs(emergency_total - float(source["emergency_kwh"].sum())) < 1e-3})
    return checks


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--result-dir", type=Path, default=ROOT / "result")
    parser.add_argument("--report", type=Path, default=ROOT / "validation" / "submission_excel_checks.json")
    args = parser.parse_args()
    checks = check_result1(args.result_dir / "result1.xlsx")
    for name, (source, adjusted) in CASES.items():
        checks.extend(check_annual(args.result_dir / name, source, adjusted))
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    failures = [item for item in checks if not item["pass"]]
    if failures:
        raise SystemExit(f"Validation failed: {failures}")
    print(f"PASS: {len(checks)} submission-workbook checks")


if __name__ == "__main__":
    main()
