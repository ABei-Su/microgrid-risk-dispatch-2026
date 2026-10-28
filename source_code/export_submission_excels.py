"""Regenerate the five official-format result workbooks from frozen outputs.

The paper reports five-seed means.  The official result workbooks require one
complete executable strategy, so they use the archived seed-1 realization.
Blank official templates are stored in ``data/templates``.  The default output
is ``reproduced/submission_excels`` and never overwrites the submitted copies.
"""
from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path
import shutil

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "data" / "templates"
INTERVAL_FILES = {
    "result2.xlsx": ("q2_interval_results.csv", False),
    "result3.xlsx": ("q3_interval_results.csv", True),
    "result4-2.xlsx": ("q4_2_interval_results.csv", False),
    "result4-3.xlsx": ("q4_3_interval_results.csv", True),
}


def _replace_rows(ws, count: int, columns: int) -> None:
    source = 2
    styles = [copy(ws.cell(source, col)._style) for col in range(1, columns + 1)]
    height = ws.row_dimensions[source].height
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for row in range(2, count + 2):
        ws.row_dimensions[row].height = height
        for col, style in enumerate(styles, 1):
            ws.cell(row, col)._style = copy(style)


def _period(first: int, last: int) -> str:
    def stamp(index: int) -> str:
        minutes = (index + 1) * 10
        suffix = "+1" if minutes >= 1440 else ""
        minutes %= 1440
        return f"{minutes // 60}:{minutes % 60:02d}{suffix}"

    return f"{stamp(first)}-{stamp(last + 1)}"


def _write_plan(ws, interval: pd.DataFrame, column: str, adjusted: bool) -> None:
    for row, (date, day) in enumerate(interval.groupby("date", sort=True), 2):
        day = day.sort_values("interval_index")
        values = day[column].to_numpy(float)
        ws.cell(row, 1, pd.Timestamp(date).to_pydatetime())
        for col, value in enumerate(values, 2):
            ws.cell(row, col, round(float(value), 6))
        ws.cell(row, 146, round(float(values.sum()), 6))
        price = day["price_yuan_per_kwh"].to_numpy(float)
        if adjusted:
            base = day["base_grid_kwh"].to_numpy(float)
            effective = day["effective_grid_kwh"].to_numpy(float)
            fee = price @ base + 1.5 * price @ np.maximum(effective - base, 0) - 0.5 * price @ np.maximum(base - effective, 0)
        else:
            fee = price @ values
        ws.cell(row, 147, round(float(fee), 6))


def _write_charge(ws, interval: pd.DataFrame) -> None:
    days = list(interval.groupby("date", sort=True))
    labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    _replace_rows(ws, len(days) * 6, 6)
    row = 2
    for date, day in days:
        day = day.sort_values("interval_index")
        for block, label in enumerate(labels):
            part = day.iloc[block * 24 : (block + 1) * 24]
            ws.cell(row, 1, pd.Timestamp(date).to_pydatetime() if block == 0 else None)
            ws.cell(row, 2, label)
            ws.cell(row, 3, round(float(part["charge_kwh"].sum()), 6))
            ws.cell(row, 4, round(float(part["discharge_kwh"].sum()), 6))
            ws.cell(row, 5, "00:00:00" if block == 0 else ("24:00" if block == 1 else None))
            ws.cell(row, 6, round(float(day.iloc[0]["soc_start_kwh"]), 6) if block == 0 else (round(float(day.iloc[-1]["soc_end_kwh"]), 6) if block == 1 else None))
            row += 1


def _write_emergency(ws, interval: pd.DataFrame) -> None:
    records: list[tuple[pd.Timestamp, str, float]] = []
    for date, day in interval.groupby("date", sort=True):
        day = day.sort_values("interval_index")
        active = day["emergency_kwh"].to_numpy(float) > 1e-7
        starts = np.flatnonzero(active & ~np.r_[False, active[:-1]])
        ends = np.flatnonzero(active & ~np.r_[active[1:], False])
        for start, end in zip(starts, ends):
            amount = float(day.iloc[start : end + 1]["emergency_kwh"].sum())
            records.append((pd.Timestamp(date), _period(int(start), int(end)), amount))
    _replace_rows(ws, max(1, len(records)), 3)
    if not records:
        ws.cell(2, 1, "无紧急购电"); ws.cell(2, 3, 0.0)
        return
    previous = None
    for row, (date, label, amount) in enumerate(records, 2):
        ws.cell(row, 1, date.to_pydatetime() if date != previous else None)
        ws.cell(row, 2, label)
        ws.cell(row, 3, round(amount, 6))
        previous = date


def _write_result1(destination: Path) -> None:
    frame = pd.read_csv(ROOT / "result" / "Q1" / "Q1_final_trajectories.csv")
    frame = frame.loc[frame["seed"].eq(1)].sort_values("slot")
    if len(frame) != 144:
        raise ValueError(f"Q1 requires 144 seed-1 intervals, found {len(frame)}")
    shutil.copy2(TEMPLATES / "result1.xlsx", destination)
    workbook = load_workbook(destination)
    plan, charge = workbook.worksheets[:2]
    for row, value in enumerate(frame["grid_kwh"], 2):
        plan.cell(row, 2, round(float(value), 6))
    for block in range(6):
        part = frame.iloc[block * 24 : (block + 1) * 24]
        charge.cell(block + 2, 2, round(float(part["charge_kwh"].sum()), 6))
        charge.cell(block + 2, 3, round(float(part["discharge_kwh"].sum()), 6))
    charge.cell(2, 5, round(float(frame.iloc[0]["soc_start_kwh"]), 6))
    charge.cell(3, 5, round(float(frame.iloc[-1]["soc_end_kwh"]), 6))
    workbook.save(destination)


def _write_annual(name: str, source: str, adjusted: bool, destination: Path) -> None:
    interval = pd.read_csv(ROOT / "figure" / "data" / source)
    counts = interval.groupby("date").size()
    if len(interval) != 334 * 144 or len(counts) != 334 or not counts.eq(144).all():
        raise ValueError(f"{name} does not contain 334 complete days")
    shutil.copy2(TEMPLATES / name, destination)
    workbook = load_workbook(destination)
    _write_plan(workbook["计划购电量"], interval, "base_grid_kwh", False)
    if adjusted:
        _write_plan(workbook["调整购电量"], interval, "effective_grid_kwh", True)
    _write_charge(workbook["充放电量"], interval)
    _write_emergency(workbook["紧急购电量"], interval)
    workbook.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=ROOT / "reproduced" / "submission_excels")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    _write_result1(args.output / "result1.xlsx")
    for name, (source, adjusted) in INTERVAL_FILES.items():
        _write_annual(name, source, adjusted, args.output / name)
    print(f"Generated five official-format workbooks in {args.output}")


if __name__ == "__main__":
    main()
